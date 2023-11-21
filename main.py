import os
import re
from tkinter import filedialog
from tkinter import *
from tkinter.ttk import Progressbar
from tkinter import messagebox

from docx import Document

selected_files = []

def select_files():
    global selected_files
    selected_files = filedialog.askopenfilenames(filetypes=[("Word 文档", "*.docx")])
    num_files = len(selected_files)
    files_label.config(text=f"已选择文件：{num_files} 个", fg="red")

def process_files():

    def close_message_box():
        message_box.destroy()
        root.deiconify()

    row = int(row_entry.get())
    column = int(column_entry.get())
    table_index = int(table_entry.get()) - 1  # 从0开始计数
    regex = regex_entry.get()

    # 隐藏窗口
    root.withdraw()

    # 创建消息框
    message_box = Toplevel(root)
    message_box.title("提示")

    # 获取屏幕宽度和高度
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()

    # 计算消息框的位置使其居中
    message_box_width = 200
    message_box_height = 100
    message_box_x = int((screen_width / 2) - (message_box_width / 2))
    message_box_y = int((screen_height / 2) - (message_box_height / 2))
    message_box.geometry(f"{message_box_width}x{message_box_height}+{message_box_x}+{message_box_y}")

    # 创建标签
    label = Label(message_box, text="转换完成")
    label.pack(pady=20)

    # 创建按钮
    button = Button(message_box, text="确定", command=close_message_box)
    button.pack(pady=10)

    # 遍历所选文件
    for i, file_path in enumerate(selected_files):
        # 打开Word文档
        doc = Document(file_path)

        # 获取指定索引的表格
        table = doc.tables[table_index]  # 根据表格索引获取表格

        # 检查指定行是否存在
        if row <= len(table.rows):
            # 获取指定行列的单元格
            cell = table.cell(row - 1, column - 1)

            # 获取单元格的文本内容
            cell_text = cell.text

            # 根据需要进行正则匹配
            if regex:
                match = re.search(regex, cell_text)
                if match:
                    cell_text = match.group(0)  # 获取第一个匹配项

            # 修改文件名
            base_path = os.path.dirname(file_path)
            new_file_name = os.path.join(base_path, cell_text + '.docx')
            os.rename(file_path, new_file_name)

            # 保存修改后的文件
            doc.save(new_file_name)

            # 输出处理后的文件路径
            print(f"已处理文件：{new_file_name}")

        # 更新进度条
        progress_bar["value"] = i + 1
        root.update_idletasks()

    # 清空选择的文件列表和进度条
    selected_files.clear()
    files_label.config(text="已选择文件：0 个", fg="red")
    progress_bar["value"] = 0

root = Tk()
root.title("DOC文档批量重命名")
root_width = 400
root_height = 300
x = int((root.winfo_screenwidth() / 2) - (root_width / 2))
y = int((root.winfo_screenheight() / 2) - (root_height / 2))
root.geometry(f"{root_width}x{root_height}+{x}+{y}")
root.resizable(False, False)
root.attributes('-topmost', True)  # 窗口置顶显示

# 创建选择文件按钮
select_files_btn = Button(root, text="选择文件", command=select_files)
select_files_btn.pack()

# 显示已选择文件数量
files_label = Label(root, text="已选择文件：0 个", fg="red")
files_label.pack()

# 创建表格标签和输入框
table_frame = Frame(root)
table_frame.pack(pady=10)

table_label = Label(table_frame, text="指定表格：")
table_label.grid(row=0, column=0, padx=5, pady=5)
table_entry = Entry(table_frame)
table_entry.grid(row=0, column=1, padx=5, pady=5)

row_label = Label(table_frame, text="指定行：")
row_label.grid(row=1, column=0, padx=5, pady=5)
row_entry = Entry(table_frame)
row_entry.grid(row=1, column=1, padx=5, pady=5)

column_label = Label(table_frame, text="指定列：")
column_label.grid(row=2, column=0, padx=5, pady=5)
column_entry = Entry(table_frame)
column_entry.grid(row=2, column=1, padx=5, pady=5)

regex_label = Label(table_frame, text="正则表达式：（选填）")
regex_label.grid(row=3, column=0, padx=5, pady=5)
regex_entry = Entry(table_frame)
regex_entry.grid(row=3, column=1, padx=5, pady=5)

# 创建批量重命名按钮
rename_btn = Button(root, text="批量重命名", command=process_files)
rename_btn.pack()

# 创建进度条
progress_bar = Progressbar(root, orient=HORIZONTAL, length=300, mode='determinate')
progress_bar.pack(pady=10)

root.mainloop()
